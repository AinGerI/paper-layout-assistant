import sys
import tempfile
import unittest
from pathlib import Path
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT/'scripts'))
from markdown_to_word import convert, parse
from audit_layout import audit

class LayoutTests(unittest.TestCase):
    def test_roles_and_override_detection(self):
        with tempfile.TemporaryDirectory() as tmp:
            d = Path(tmp); src=d/'in.md'; out=d/'out.docx'; profile=ROOT/'profiles/chinese-paper.json'
            src.write_text('# 题目\n\n## 摘要\n\n摘要正文。\n\n关键词：测试\n\n## 第一章 引言\n\n正文 $CO_2$ & AP。\n\n## 参考文献\n\n文献。')
            convert(src,out,d/'reports',profile)
            self.assertTrue(audit(out,profile)['passed'])
            doc=Document(out)
            self.assertEqual(doc.styles['Heading 1'].paragraph_format.alignment, WD_ALIGN_PARAGRAPH.CENTER)
            self.assertEqual(doc.styles['PWD Front Heading'].paragraph_format.alignment, WD_ALIGN_PARAGRAPH.CENTER)
            self.assertEqual([p.style.name for p in doc.paragraphs], ['Title','PWD Front Heading','Normal','PWD Keywords','Heading 1','Normal','PWD Front Heading','PWD References'])
            doc.paragraphs[4].alignment=WD_ALIGN_PARAGRAPH.LEFT
            doc.paragraphs[5]._p.get_or_add_pPr().get_or_add_ind().set(qn('w:firstLineChars'),'0')
            doc.save(out)
            result=audit(out,profile)
            self.assertFalse(result['passed'])
            self.assertEqual(len(result['errors']),2)
    def test_literal_ampersand_guard(self):
        with self.assertRaisesRegex(ValueError,'connecting text'):
            parse(r'正文 $CO_2\&\mathrm{AP}$')
        parse(r'正文 $CO_2$ & AP')
        parse(r'\begin{equation}\begin{aligned}x&=1\\y&=2\end{aligned}\end{equation}')

    def test_caption_spacing_overrides(self):
        with tempfile.TemporaryDirectory() as tmp:
            d=Path(tmp); profile=ROOT/'profiles/chinese-paper.json'
            (d/'in.md').write_text('# 题目\n\n表1 数据\n\n<table><tr><td>数值</td></tr><tr><td>1</td></tr></table>')
            convert(d/'in.md',d/'out.docx',d/'reports',profile)
            self.assertTrue(audit(d/'out.docx',profile)['passed'])
            for prop, value in [('after','600'), ('before','0'), ('afterAutospacing','1'), ('beforeLines','100')]:
                doc=Document(d/'out.docx')
                caption=doc.paragraphs[1]
                caption._p.get_or_add_pPr().get_or_add_spacing().set(qn('w:'+prop),value)
                doc.save(d/'bad.docx')
                result=audit(d/'bad.docx',profile)
                self.assertFalse(result['passed'],prop)
                self.assertTrue(any(e.get('property')=='spacing.'+prop for e in result['errors']),prop)
