"""Strict manuscript Markdown subset -> visible-TeX Word with PWD_EQ bookmarks."""
from pathlib import Path
from html.parser import HTMLParser
import argparse
import html
import io
import json
import math
import re
from docx import Document
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from PIL import Image

ROLES = {'title': 'Title', 'heading_1': 'Heading 1', 'heading_2': 'Heading 2',
         'heading_3': 'Heading 3', 'body': 'Normal', 'caption': 'PWD Caption',
         'equation': 'PWD Formula Input', 'table': 'PWD Table Text', 'references': 'PWD References',
         'front_heading': 'PWD Front Heading', 'keywords': 'PWD Keywords'}
ALIGN = {'left': WD_ALIGN_PARAGRAPH.LEFT, 'right': WD_ALIGN_PARAGRAPH.RIGHT,
         'center': WD_ALIGN_PARAGRAPH.CENTER, 'justify': WD_ALIGN_PARAGRAPH.JUSTIFY}
ROLE_KEYS = {'east_asia', 'latin', 'size_pt', 'bold', 'alignment', 'line_spacing',
             'before_pt', 'after_pt', 'first_line_chars'}

def require(condition, message):
    if not condition:
        raise ValueError(message)

def numeric(value, label, minimum=0, strict=False):
    require(type(value) in (int, float) and math.isfinite(value) and
            (value > minimum if strict else value >= minimum), f'Invalid {label}: {value!r}')

def keys(obj, expected, label):
    require(isinstance(obj, dict) and set(obj) == set(expected), f'{label} keys must be {sorted(expected)}')

def load_profile(path):
    p = json.loads(path.read_text(encoding='utf-8'))
    keys(p, {'schema_version', 'name', 'page', 'fonts', 'roles', 'figure', 'table', 'equations'}, 'profile')
    require(p['schema_version'] == 1, 'Unsupported schema_version')
    require(isinstance(p['name'], str), 'Profile name must be text')
    keys(p['page'], {'width_cm', 'height_cm', 'top_cm', 'bottom_cm', 'left_cm', 'right_cm'}, 'page')
    for k, v in p['page'].items(): numeric(v, k, strict=k in ('width_cm', 'height_cm'))
    require(p['page']['width_cm'] > p['page']['left_cm'] + p['page']['right_cm'], 'Margins consume page width')
    require(p['page']['height_cm'] > p['page']['top_cm'] + p['page']['bottom_cm'], 'Margins consume page height')
    keys(p['fonts'], {'east_asia', 'latin'}, 'fonts')
    for v in p['fonts'].values(): require(isinstance(v, str) and v.strip(), 'Font names must be nonempty')
    p['roles'].setdefault('front_heading', dict(p['roles']['heading_1'], alignment='center', first_line_chars=0))
    p['roles'].setdefault('keywords', dict(first_line_chars=0, alignment='left'))
    keys(p['roles'], ROLES, 'roles')
    body_required = ROLE_KEYS - {'east_asia', 'latin'}
    require(body_required <= set(p['roles']['body']), 'body must specify all non-font properties')
    resolved = {}
    for role, values in p['roles'].items():
        require(isinstance(values, dict) and set(values) <= ROLE_KEYS, f'Unknown property in role {role}')
        r = dict(p['fonts'], **p['roles']['body']); r.update(values)
        for k in ('east_asia', 'latin'): require(isinstance(r[k], str) and r[k].strip(), f'Invalid font in {role}')
        for k in ('size_pt', 'line_spacing'): numeric(r[k], f'{role}.{k}', strict=True)
        for k in ('before_pt', 'after_pt', 'first_line_chars'): numeric(r[k], f'{role}.{k}')
        require(type(r['bold']) is bool and r['alignment'] in ALIGN, f'Invalid bold/alignment in {role}')
        resolved[role] = r
    keys(p['figure'], {'max_width_cm', 'max_height_cm'}, 'figure')
    for k, v in p['figure'].items(): numeric(v, k, strict=True)
    keys(p['table'], {'border', 'header_repeat', 'allow_row_split'}, 'table')
    require(p['table']['border'] in ('three_line', 'grid', 'none'), 'table.border must be three_line, grid or none')
    for k in ('header_repeat', 'allow_row_split'): require(type(p['table'][k]) is bool, f'table.{k} must be boolean')
    keys(p['equations'], {'disable_grid', 'single_omega_tilde_compat'}, 'equations')
    for k, v in p['equations'].items(): require(type(v) is bool, f'equations.{k} must be boolean')
    return p, resolved

class TableParser(HTMLParser):
    def __init__(self):
        super().__init__(); self.rows = []; self.cell = None; self.stack = []
    def handle_starttag(self, tag, attrs):
        require(tag in ('table', 'thead', 'tbody', 'tr', 'td', 'th', 'br'), f'Unsupported table tag: {tag}')
        attrs = dict(attrs)
        require(set(attrs) <= {'rowspan', 'colspan'} if tag in ('td', 'th') else not attrs, f'Unsupported table attributes: {attrs}')
        require(all(str(v) == '1' for v in attrs.values()), 'Merged table cells are not supported')
        if tag == 'br':
            require(self.cell is not None, 'br must be inside a table cell'); self.cell.append('\n'); return
        parent = self.stack[-1] if self.stack else None
        allowed = {'table': (None,), 'thead': ('table',), 'tbody': ('table',), 'tr': ('table', 'thead', 'tbody'), 'td': ('tr',), 'th': ('tr',)}
        require(parent in allowed[tag], f'Malformed table nesting at {tag}')
        self.stack.append(tag)
        if tag == 'tr': self.rows.append([])
        if tag in ('td', 'th'):
            self.cell = []; self.rows[-1].append(self.cell)
    def handle_startendtag(self, tag, attrs):
        require(tag == 'br', 'Only br may be self-closing'); self.handle_starttag(tag, attrs)
    def handle_data(self, text):
        if self.cell is not None: self.cell.append(text)
        else: require(not text.strip(), 'Text outside table cells')
    def handle_endtag(self, tag):
        require(self.stack and self.stack[-1] == tag, f'Malformed table closing tag: {tag}')
        self.stack.pop()
        if tag in ('td', 'th'): self.cell = None
    def result(self):
        require(not self.stack and self.rows and self.rows[0], 'Incomplete or empty table')
        rows = [[''.join(c) for c in row] for row in self.rows]
        require(all(len(row) == len(rows[0]) for row in rows), 'Table must be rectangular')
        return rows

INLINE = re.compile(r'\\\(([\s\S]*?)\\\)|(?<!\\)\$([^$\n]+?)(?<!\\)\$')

def plain_check(text):
    for match in INLINE.finditer(text):
        body = match[1] if match[1] is not None else match[2]
        require(r'\&' not in body, 'Literal \u0026 inside formula: move connecting text outside formula delimiters (e.g. $CO_2$ & AP); do not change array alignment markers')
        require(body.strip() and not re.search(r'\\[()\[\]]|(?<!\\)\$|\\(?:begin|end)\{equation', body), 'Empty or nested inline formula')
    remainder = INLINE.sub('', text)
    require(not re.search(r'\\[()\[\]]|\\(?:begin|end)\{|(?<!\\)\$', remainder), 'Unclosed or unsupported formula delimiters')
    require(not re.search(r'`|\*\*|__|\[\^|\[[^\]]+\]\([^)]*\)|<[/a-zA-Z!]', remainder), 'Unsupported inline Markdown/HTML: use plain text and formula delimiters')
    require(not re.search(r'(?<!\\)\*[^*\n]+(?<!\\)\*', remainder), 'Markdown emphasis is not supported')

def parse(text):
    blocks = []
    for b in re.split(r'\n\s*\n', text.strip()):
        b = b.strip()
        if not b: continue
        if b.startswith('<table'):
            t = TableParser(); t.feed(b); t.close(); blocks.append(dict(type='table', rows=t.result()))
        elif b.startswith('!['):
            m = re.fullmatch(r'!\[(.*?)\]\((.*?)\)', b); require(m is not None, 'Image must be standalone ![alt](relative/path)')
            blocks.append(dict(type='image', alt=m[1], path=m[2]))
        elif b.startswith('#'):
            m = re.fullmatch(r'(#{1,4})\s+([^\n]+)', b); require(m is not None, 'Use headings # through #### on their own line')
            blocks.append(dict(type='heading', level=len(m[1]), text=m[2]))
        elif b.startswith((r'\begin{equation}', r'\begin{equation*}', r'\[')):
            start = r'\[' if b.startswith(r'\[') else re.match(r'\\begin\{equation\*?\}', b)[0]
            end = r'\]' if start == r'\[' else start.replace('begin', 'end')
            require(b.endswith(end), 'Incomplete display equation')
            body = b[len(start):-len(end)].strip(); number = 'auto' if start == r'\begin{equation}' else None
            tag = re.search(r'\\tag\{([^}]+)\}\s*$', body)
            if tag and start == r'\[': body = body[:tag.start()].strip(); number = tag[1]
            require(body and not re.search(r'\\(?:tag|label)\b|\\(?:begin|end)\{equation\*?\}|\\[\[\]]', body), 'Empty/nested equation or unsupported manual numbering')
            blocks.append(dict(type='display', body=' '.join(body.splitlines()), number=number))
        else:
            require(not re.search(r'^\s*(?:[-*+]\s|\d+[.)]\s|>|\||```|~~~|---\s*$)', b, re.M), 'Lists, pipe tables, quotes, front matter and code fences are not supported')
            require(not re.search(r'^\s*#', b, re.M), 'Separate headings and blocks with a blank line')
            blocks.append(dict(type='paragraph', text=html.unescape(b.replace('\n', ' '))))
    for b in blocks:
        if 'text' in b: plain_check(b['text'])
        for row in b.get('rows', []):
            for c in row: plain_check(c)
    require(blocks, 'Markdown is empty')
    return blocks

def xml_property(parent, name, value):
    node = parent.find(qn('w:' + name))
    if node is None: node = OxmlElement('w:' + name); parent.append(node)
    node.set(qn('w:val'), str(value))

def convert(source, target, work, profile_path):
    require(not target.exists(), f'Output already exists: {target}; choose a new filename')
    require(source.resolve() != target.resolve(), 'Input and output must differ')
    p, roles = load_profile(profile_path); blocks = parse(source.read_text(encoding='utf-8-sig'))
    require(not any((work / n).exists() for n in ('blocks.json', 'equations.json', 'coverage.json', 'profile.json')), 'Work directory already contains a conversion report; choose a new directory')
    doc = Document(); sec = doc.sections[0]; page = p['page']
    sec.page_width = Cm(page['width_cm']); sec.page_height = Cm(page['height_cm'])
    for side in ('top', 'bottom', 'left', 'right'): setattr(sec, side + '_margin', Cm(page[side + '_cm']))
    for role, name in ROLES.items():
        st = doc.styles[name] if name in doc.styles else doc.styles.add_style(name, 1)
        if role != 'body': st.base_style = doc.styles['Normal']
        # The bundled Word template contains theme colors, title borders and tracking.
        # Reset role properties so only this profile controls the generated document.
        for properties in (st.element.get_or_add_rPr(), st.element.get_or_add_pPr()):
            for child in list(properties): properties.remove(child)
        linked = st.element.find(qn('w:link'))
        if linked is not None: st.element.remove(linked)
        r = roles[role]; st.font.name = r['latin']; st.font.size = Pt(r['size_pt']); st.font.bold = r['bold']
        st.element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'), r['east_asia'])
        xml_property(st.element.get_or_add_rPr(), 'color', '000000')
        xml_property(st.element.get_or_add_rPr(), 'szCs', round(r['size_pt'] * 2))
        xml_property(st.element.get_or_add_rPr(), 'bCs', int(r['bold']))
        pf = st.paragraph_format; pf.alignment = ALIGN[r['alignment']]; pf.line_spacing = r['line_spacing']
        pf.space_before = Pt(r['before_pt']); pf.space_after = Pt(r['after_pt']); pf.first_line_indent = Pt(r['first_line_chars'] * r['size_pt'])
        st.element.get_or_add_pPr().get_or_add_ind().set(qn('w:firstLineChars'), str(round(r['first_line_chars'] * 100)))
        pf.keep_with_next = role == 'title' or role.startswith('heading_'); pf.keep_together = False
        pf.page_break_before = False; pf.widow_control = True
        st.next_paragraph_style = doc.styles['Normal']
        xml_property(st.element.get_or_add_pPr(), 'outlineLvl', int(role[-1])-1 if role.startswith('heading_') else 9)
    image_style = doc.styles.add_style('PWD Image', 1); image_style.base_style = doc.styles['PWD Caption']
    image_style.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    equations = []; coverage = []; captions = []; changes = []; pictures = []
    def equation(par, code, kind, block):
        if p['equations']['single_omega_tilde_compat']:
            code, count = re.subn(r'\\widetilde\{\s*\\omega\s*\}', lambda m: m[0].replace('widetilde', 'tilde'), code)
            if count: changes.append(dict(source_block=block, rule='single_omega_tilde', occurrences=count))
        n = len(equations) + 1; name = f'PWD_EQ_{n:06d}'
        start = OxmlElement('w:bookmarkStart'); start.set(qn('w:id'), str(n)); start.set(qn('w:name'), name)
        end = OxmlElement('w:bookmarkEnd'); end.set(qn('w:id'), str(n))
        par._p.append(start); par.add_run(code); par._p.append(end)
        if p['equations']['disable_grid']:
            for prop in ('snapToGrid', 'adjustRightInd'): xml_property(par._p.get_or_add_pPr(), prop, 0)
        equations.append(dict(bookmark=name, kind=kind, code=code, source_block=block))
    def add(par, text, block):
        pos = 0
        for m in INLINE.finditer(text):
            require((m[1] if m[1] is not None else m[2]).strip(), 'Empty inline formula')
            par.add_run(text[pos:m.start()].replace(r'\$', '$')); equation(par, '$' + (m[1] if m[1] is not None else m[2]).strip() + '$', 'inline', block); pos = m.end()
        par.add_run(text[pos:].replace(r'\$', '$'))
    def borders(cell, top, bottom):
        el = OxmlElement('w:tcBorders')
        for edge in ('top', 'left', 'bottom', 'right'):
            enabled = p['table']['border'] == 'grid' or p['table']['border'] == 'three_line' and ((edge == 'top' and top) or (edge == 'bottom' and bottom))
            node = OxmlElement('w:' + edge); node.set(qn('w:val'), 'single' if enabled else 'nil')
            if enabled: node.set(qn('w:sz'), '8'); node.set(qn('w:color'), '000000')
            el.append(node)
        cell._tc.get_or_add_tcPr().append(el)
    ref = False
    for bi, b in enumerate(blocks):
        kind = b['type']; prev = blocks[bi-1]['type'] if bi else None; nxt = blocks[bi+1]['type'] if bi+1 < len(blocks) else None
        if kind == 'heading':
            ref = b['text'].strip() == '参考文献'
            style = 'PWD Front Heading' if b['text'].strip() in ('摘要', '参考文献') else 'PWD Keywords' if re.match(r'^关键词[：:]', b['text']) else 'Title' if b['level'] == 1 else f"Heading {b['level']-1}"
            par = doc.add_paragraph(style=style); add(par, b['text'], bi)
        elif kind == 'paragraph':
            ahead_table = nxt == 'table' or bi+2 < len(blocks) and blocks[bi+2]['type'] == 'table' and blocks[bi+1].get('text', '').startswith('(')
            caption = bool(ahead_table and re.match(r'^表\s*\d+', b['text']) or (prev == 'image' or nxt == 'image') and re.match(r'^图\s*\d+', b['text']) or nxt == 'table' and b['text'].startswith('(') and bi-1 in captions)
            par = doc.add_paragraph(style='PWD Keywords' if re.match(r'^关键词[：:]', b['text']) else 'PWD Caption' if caption else 'PWD References' if ref else 'Normal'); add(par, b['text'].lstrip() if ref else b['text'], bi)
            if caption:
                captions.append(bi); par.paragraph_format.keep_with_next = nxt in ('table', 'image') or ahead_table
        elif kind == 'display':
            par = doc.add_paragraph(style='PWD Formula Input'); numbered = b['number'] is not None
            equation(par, r'\begin{equation}' + b['body'] + r'\end{equation}' if numbered else r'\[' + b['body'] + r'\]', 'numbered' if numbered else 'display', bi)
            equations[-1]['source_number'] = b['number']
        elif kind == 'image':
            require(not re.match(r'^[a-zA-Z]+://', b['path']), 'Images must use local relative paths')
            path = source.parent / b['path']
            with Image.open(path) as im:
                width, height = im.size
                raw_dpi = im.info.get('dpi', ())
                valid_dpi = isinstance(raw_dpi, (tuple, list)) and len(raw_dpi) >= 2 and all(type(v) in (int, float) and math.isfinite(v) and v > 0 for v in raw_dpi[:2])
                # A single scale preserves pixel aspect ratio, including unequal DPI metadata.
                dpi = max(raw_dpi[:2]) if valid_dpi else 96.0
                stream = io.BytesIO(); im.save(stream, format='PNG'); stream.seek(0)
            maxw = min(p['figure']['max_width_cm'], page['width_cm']-page['left_cm']-page['right_cm'])
            maxh = min(p['figure']['max_height_cm'], page['height_cm']-page['top_cm']-page['bottom_cm'])
            scale = min(2.54/dpi, maxw/width, maxh/height)
            pictures.append(dict(source_block=bi, path=b['path'], pixels=[width, height],
                                 source_dpi=list(raw_dpi[:2]) if valid_dpi else repr(raw_dpi), dpi_used=dpi,
                                 sizing_basis='metadata_max_axis' if valid_dpi else 'estimated_96dpi',
                                 width_cm=width*scale, height_cm=height*scale))
            par = doc.add_paragraph(style='PWD Image'); par.add_run().add_picture(stream, width=Cm(width*scale))
            if bi+1 < len(blocks) and re.match(r'^图\s*\d+', blocks[bi+1].get('text', '')): par.paragraph_format.keep_with_next = True
        elif kind == 'table':
            rows = b['rows']; cols = len(rows[0]); table = doc.add_table(rows=len(rows), cols=cols); table.autofit = False
            width = (page['width_cm']-page['left_cm']-page['right_cm'])/cols
            for col in table.columns: col.width = Cm(width)
            for ri, row in enumerate(rows):
                trpr = table.rows[ri]._tr.get_or_add_trPr()
                if not p['table']['allow_row_split']: trpr.append(OxmlElement('w:cantSplit'))
                if ri == 0 and p['table']['header_repeat']: trpr.append(OxmlElement('w:tblHeader'))
                for ci, text in enumerate(row):
                    cell = table.cell(ri, ci); cell.width = Cm(width); par = cell.paragraphs[0]; par.style = 'PWD Table Text'; add(par, text, bi)
                    borders(cell, ri == 0, ri in (0, len(rows)-1))
                    if ri == 0:
                        for run in par.runs: run.bold = True
        coverage.append(dict(source_block=bi, type=kind))
    footer = sec.footer.paragraphs[0]; footer.style = 'PWD Caption'
    field = OxmlElement('w:fldSimple'); field.set(qn('w:instr'), 'PAGE'); footer._p.append(field)
    target.parent.mkdir(parents=True, exist_ok=True); work.mkdir(parents=True, exist_ok=True)
    with target.open('xb') as stream: doc.save(stream)
    reports = {'blocks.json': blocks, 'equations.json': equations, 'profile.json': p,
               'coverage.json': dict(blocks=coverage, captions=captions, compatibility_changes=changes, pictures=pictures,
                                    formulas=len(equations), tables=sum(b['type']=='table' for b in blocks), images=sum(b['type']=='image' for b in blocks))}
    for name, data in reports.items():
        with (work/name).open('x', encoding='utf-8') as f: json.dump(data, f, ensure_ascii=False, indent=2)
    return reports['coverage.json']

def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('input', type=Path); parser.add_argument('output', type=Path)
    parser.add_argument('--work-dir', type=Path, required=True)
    parser.add_argument('--profile', type=Path, default=Path(__file__).resolve().parents[1]/'profiles/chinese-paper.json')
    args = parser.parse_args()
    try: result = convert(args.input, args.output, args.work_dir, args.profile)
    except (ValueError, OSError, KeyError) as exc: parser.exit(2, f'Conversion failed: {exc}\n')
    print(json.dumps(dict(output=str(args.output), **{k: result[k] for k in ('formulas', 'tables', 'images')}), ensure_ascii=False))

if __name__ == '__main__': main()
